"""ExportService: serializes VERIFIED Leads for external consumption.

Unlike collectors and production scraping, exporting already-persisted,
already-verified Leads involves no external data collection and no risk
of fabrication -- it is a pure serialization of data already in the
database, so the foundation implements it fully rather than as a
placeholder.

Only ``VerificationStatus.VERIFIED`` Leads are ever exported, in every
format -- this is an unconditional platform rule, not a default that a
caller can override. Raw Signals are never exported; a Lead's "Source"
column is a read-only, human-readable summary (the distinct collector
source names behind its supporting signals), not the signals themselves.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import UTC, datetime
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.application.interfaces.repositories import LeadRepository, SignalRepository
from app.config.settings import ExcelExportConfig
from app.core.exceptions import ExportError
from app.domain.enums import ServiceCategory, VerificationStatus
from app.domain.lead import Lead

ExportFormat = Literal["json", "csv", "xlsx"]

#: Column order and display headers shared by every tabular export format.
_COLUMNS: tuple[str, ...] = (
    "Lead ID",
    "Lead Type",
    "Intent Score",
    "Confidence Score",
    "Priority",
    "Verification Status",
    "Source",
    "Municipality",
    "County",
    "Reasoning",
    "Created At",
)

#: Columns that must be written as native numeric/datetime cell values in
#: the Excel export, rather than text -- "proper column types" per spec.
_NUMERIC_COLUMNS = frozenset({"Intent Score", "Confidence Score"})
_DATE_COLUMNS = frozenset({"Created At"})


class ExportService:
    """Serializes VERIFIED Leads currently in the LeadRepository."""

    def __init__(
        self,
        lead_repository: LeadRepository,
        signal_repository: SignalRepository,
        excel_config: ExcelExportConfig | None = None,
    ) -> None:
        self._lead_repository = lead_repository
        self._signal_repository = signal_repository
        self._excel_config = excel_config or ExcelExportConfig()

    def export_leads(
        self,
        *,
        export_format: ExportFormat,
        lead_type: ServiceCategory | None = None,
        limit: int = 1000,
    ) -> str | bytes:
        """Return VERIFIED leads matching ``lead_type`` serialized as ``export_format``.

        Returns ``str`` for ``json``/``csv`` and ``bytes`` for ``xlsx``.
        """
        leads = self._lead_repository.list_all(
            lead_type=lead_type,
            verification_status=VerificationStatus.VERIFIED,
            limit=limit,
        )
        rows = [self._row_for(lead) for lead in leads]

        if export_format == "json":
            return self._to_json(leads)
        if export_format == "csv":
            return self._to_csv(rows)
        if export_format == "xlsx":
            return self._to_xlsx(rows, self._excel_config)
        raise ExportError(f"Unsupported export format: {export_format!r}.")

    def _row_for(self, lead: Lead) -> dict[str, object]:
        return {
            "Lead ID": str(lead.id),
            "Lead Type": lead.lead_type.value,
            "Intent Score": lead.intent_score,
            "Confidence Score": lead.estimated_confidence,
            "Priority": lead.priority.value,
            "Verification Status": lead.verification_status.value,
            "Source": self._source_summary(lead),
            "Municipality": lead.municipality,
            "County": lead.county,
            "Reasoning": lead.reasoning,
            "Created At": lead.created_at,
        }

    def _source_summary(self, lead: Lead) -> str:
        """Return a human-readable summary of which public sources back this lead.

        Reads only the ``source`` field of each supporting signal -- never
        the signals themselves -- so raw Signal data is never exposed
        through an export.
        """
        signals = self._signal_repository.list_by_ids(lead.supporting_signal_ids)
        distinct_sources = sorted({signal.source for signal in signals})
        return "; ".join(distinct_sources) if distinct_sources else "unknown"

    @staticmethod
    def _to_json(leads: list[Lead]) -> str:
        payload = [json.loads(lead.model_dump_json()) for lead in leads]
        return json.dumps(payload, indent=2, ensure_ascii=False)

    @staticmethod
    def _to_csv(rows: list[dict[str, object]]) -> str:
        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=list(_COLUMNS))
        writer.writeheader()
        for row in rows:
            serialized = dict(row)
            created_at = serialized["Created At"]
            if isinstance(created_at, datetime):
                serialized["Created At"] = created_at.isoformat()
            writer.writerow(serialized)
        return buffer.getvalue()

    @staticmethod
    def _to_xlsx(rows: list[dict[str, object]], config: ExcelExportConfig) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Leads"

        worksheet.append(list(_COLUMNS))
        _style_header_row(worksheet, config)

        for row in rows:
            worksheet.append([_xlsx_cell_value(row[column]) for column in _COLUMNS])

        for row_cells in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for cell in row_cells:
                column_name = _COLUMNS[cell.column - 1]
                if column_name in _DATE_COLUMNS:
                    cell.number_format = config.date_format
                elif column_name in _NUMERIC_COLUMNS:
                    cell.number_format = "0.00"

        if config.enable_autofilter and worksheet.max_row >= 1:
            worksheet.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{worksheet.max_row}"
        if config.freeze_header_row:
            worksheet.freeze_panes = "A2"
        if config.auto_size_columns:
            _auto_size_columns(worksheet, config)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()


def _xlsx_cell_value(value: object) -> object:
    """Make a row value safe for an openpyxl cell.

    Excel's datetime cells have no timezone concept -- openpyxl rejects a
    timezone-aware ``datetime`` outright. Every timestamp in this platform
    is stored in UTC, so converting to UTC and dropping ``tzinfo`` loses
    no information the offset wasn't already redundant with; the column
    header and this note establish the convention.
    """
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(UTC).replace(tzinfo=None)
    return value


def _style_header_row(worksheet: Worksheet, config: ExcelExportConfig) -> None:
    fill = PatternFill(
        start_color=config.header_fill_color, end_color=config.header_fill_color, fill_type="solid"
    )
    font = Font(color=config.header_font_color, bold=True)
    alignment = Alignment(horizontal="center", vertical="center")
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment


def _auto_size_columns(worksheet: Worksheet, config: ExcelExportConfig) -> None:
    for index, column_name in enumerate(_COLUMNS, start=1):
        longest = len(column_name)
        for cell in worksheet[get_column_letter(index)][1:]:
            value = cell.value
            text_length = len(str(value)) if value is not None else 0
            longest = max(longest, text_length)
        width = min(max(longest + 2, config.min_column_width), config.max_column_width)
        worksheet.column_dimensions[get_column_letter(index)].width = width
